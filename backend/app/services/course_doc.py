"""Course design / graduation design document generator (feature 11)."""
from __future__ import annotations

import io
from datetime import datetime
from typing import Any

from app.services.schema_parser import ParsedSchema


_TEXTS: dict[str, dict[str, Any]] = {
    "zh": {
        "doc_title": "{project} 数据库设计说明书",
        "author_line": "**作者**：{author}　**学号**：{student}",
        "generated": "> 自动生成于 {now} · AI Database Architect",
        "overview": "## 1. 项目概述",
        "db_name": "- 数据库名称：`{name}`",
        "db_version": "- 数据库版本：{version}",
        "table_count": "- 实体/表数量：{n}",
        "rel_count": "- 关系数量：{n}",
        "field_count": "- 字段总数：{n}",
        "req": "## 2. 需求分析",
        "req_note": "> 说明：请根据实际业务需求补充需求描述。以下为从数据库结构中识别出的实体与联系概要。",
        "entities_found": "识别出的实体：",
        "relations_found": "识别出的联系：",
        "concept": "## 3. 概念结构设计",
        "concept_note": "采用 Chen 记法 ER 概念模型描述业务实体、属性与联系。概念模型图可在『ER 模型 → 概念模型』中查看与导出。",
        "entity_attrs": "### 3.1 实体及其属性",
        "relations_title": "### 3.2 联系",
        "no_relations": "无。",
        "pk_tag": "（主键）",
        "logical": "## 4. 逻辑结构设计",
        "logical_note": "将概念模型转换为关系模式（表结构）如下：",
        "col_field": "字段",
        "col_type": "类型",
        "col_null": "可空",
        "col_pk": "主键",
        "col_uniq": "唯一",
        "col_default": "默认值",
        "col_comment": "注释",
        "yes": "是",
        "no": "否",
        "norm": "## 5. 规范化分析",
        "norm_note": "> 提示：可结合『架构评审』中的 AI 意见补充范式分析（1NF/2NF/3NF/BCNF）章节。",
        "norm_line": "- `{table}`：主键 {pk}；满足 1NF（属性原子性）。",
        "no_pk": "无",
        "rels": "## 6. 关系说明",
        "rel_cols": ["源表.字段", "目标表.字段", "基数", "类型", "状态"],
        "dict": "## 7. 数据字典",
        "dict_note": "各表字段字典已在上文『逻辑结构设计』中列出，可在『注释补全』页面导出完整数据字典（Excel/Word/Markdown）。",
        "summary": "## 8. 总结",
        "summary_text": "本系统数据库共设计 {tables} 张表、{fields} 个字段、{rels} 条关系，覆盖了业务的核心数据需求，并遵循数据库规范化原则。",
        "colon": "：",
        "join": "、",
        "paren": "（{v}）",
        "arrow": "→",
    },
    "en": {
        "doc_title": "{project} Database Design Document",
        "author_line": "**Author**: {author}　**Student ID**: {student}",
        "generated": "> Auto-generated on {now} · AI Database Architect",
        "overview": "## 1. Project Overview",
        "db_name": "- Database name: `{name}`",
        "db_version": "- Database version: {version}",
        "table_count": "- Entities/tables: {n}",
        "rel_count": "- Relationships: {n}",
        "field_count": "- Total fields: {n}",
        "req": "## 2. Requirements Analysis",
        "req_note": "> Note: please complete the requirements description based on the actual business needs. The following is a summary of the entities and relationships identified from the database structure.",
        "entities_found": "Identified entities:",
        "relations_found": "Identified relationships:",
        "concept": "## 3. Conceptual Design",
        "concept_note": "The Chen-notation ER concept model describes business entities, attributes and relationships. The concept diagram can be viewed and exported in 'ER Model → Concept Model'.",
        "entity_attrs": "### 3.1 Entities and Attributes",
        "relations_title": "### 3.2 Relationships",
        "no_relations": "None.",
        "pk_tag": " (PK)",
        "logical": "## 4. Logical Design",
        "logical_note": "The relational schemas (table structures) converted from the concept model are as follows:",
        "col_field": "Field",
        "col_type": "Type",
        "col_null": "Nullable",
        "col_pk": "PK",
        "col_uniq": "Unique",
        "col_default": "Default",
        "col_comment": "Comment",
        "yes": "Yes",
        "no": "No",
        "norm": "## 5. Normalization Analysis",
        "norm_note": "> Tip: you can extend this section with normalization analysis (1NF/2NF/3NF/BCNF) based on the AI review in 'Architecture Review'.",
        "norm_line": "- `{table}`: PK {pk}; satisfies 1NF (atomic attributes).",
        "no_pk": "none",
        "rels": "## 6. Relationship Description",
        "rel_cols": ["Source.Field", "Target.Field", "Cardinality", "Type", "Status"],
        "dict": "## 7. Data Dictionary",
        "dict_note": "The field dictionary of every table is listed in 'Logical Design' above; you can export a full data dictionary (Excel/Word/Markdown) from the 'Comment Completion' page.",
        "summary": "## 8. Summary",
        "summary_text": "The database contains {tables} tables, {fields} fields and {rels} relationships, covering the core data needs of the business and following database normalization principles.",
        "colon": ": ",
        "join": ", ",
        "paren": " ({v})",
        "arrow": "→",
    },
}


def _texts(lang: str) -> dict[str, Any]:
    return _TEXTS.get("en" if str(lang).startswith("en") else "zh")


def _rel_display(r: dict, T: dict) -> str:
    src = f"{r.get('source_table')}.{r.get('source_column')}"
    tgt = f"{r.get('target_table')}.{r.get('target_column')}"
    return f"{src} {T['arrow']} {tgt}{T['paren'].format(v=r.get('cardinality') or '')}"


def _entity_name(entities: list[dict], eid: str) -> str:
    for e in entities:
        if e.get("id") == eid:
            return e.get("name") or eid
    return eid


def build_markdown(
    project_name: str,
    project_description: str | None,
    schema: ParsedSchema,
    relationships: list[dict],
    concept_model: dict | None = None,
    author: str = "",
    student_id: str = "",
    lang: str = "zh",
) -> str:
    T = _texts(lang)
    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    entities = concept_model.get("entities", []) if concept_model else []
    relations = concept_model.get("relations", []) if concept_model else []
    lines: list[str] = []

    lines.append(T["doc_title"].format(project=project_name))
    if author or student_id:
        lines.append("")
        lines.append(T["author_line"].format(author=author or "-", student=student_id or "-"))
    lines.append("")
    lines.append(T["generated"].format(now=now))
    lines.append("")

    # 1 概述
    lines.append(T["overview"])
    lines.append("")
    if project_description:
        lines.append(project_description)
        lines.append("")
    lines.append(T["db_name"].format(name=schema.database_name))
    if schema.db_version:
        lines.append(T["db_version"].format(version=schema.db_version))
    lines.append(T["table_count"].format(n=len(schema.tables)))
    lines.append(T["rel_count"].format(n=len(relationships)))
    lines.append(T["field_count"].format(n=sum(len(t.columns) for t in schema.tables)))
    lines.append("")

    # 2 需求分析
    lines.append(T["req"])
    lines.append("")
    lines.append(T["req_note"])
    lines.append("")
    if entities:
        lines.append(T["entities_found"])
        lines.append("")
        for e in entities:
            attrs = T["join"].join(a["name"] for a in (e.get("attributes") or [])[:8])
            lines.append(f"- **{e['name']}**{T['colon']}{attrs}")
        lines.append("")
    if relations:
        lines.append(T["relations_found"])
        lines.append("")
        for r in relations:
            src = _entity_name(entities, r.get("source_entity") or "")
            tgt = _entity_name(entities, r.get("target_entity") or "")
            lines.append(
                f"- {r['name']}{T['colon']}{src}{T['paren'].format(v=r.get('source_card'))} "
                f"— {tgt}{T['paren'].format(v=r.get('target_card'))}"
            )
        lines.append("")

    # 3 概念结构设计
    lines.append(T["concept"])
    lines.append("")
    lines.append(T["concept_note"])
    lines.append("")
    lines.append(T["entity_attrs"])
    lines.append("")
    for e in entities:
        lines.append(f"#### {e['name']}")
        lines.append("")
        for a in e.get("attributes") or []:
            tag = T["pk_tag"] if a.get("is_pk") else ""
            lines.append(f"- {a['name']}{tag}")
        lines.append("")
    lines.append(T["relations_title"])
    lines.append("")
    if relations:
        lines.append(f"| {T['col_field']} | Entity A | {T['col_type']} | Entity B | {T['col_type']} |")
        lines.append("|------|----------|------|----------|------|")
        for r in relations:
            src = _entity_name(entities, r.get("source_entity") or "")
            tgt = _entity_name(entities, r.get("target_entity") or "")
            lines.append(
                f"| {r['name']} | {src} | {r.get('source_card')} | "
                f"{tgt} | {r.get('target_card')} |"
            )
    else:
        lines.append(T["no_relations"])
    lines.append("")

    # 4 逻辑结构设计
    lines.append(T["logical"])
    lines.append("")
    lines.append(T["logical_note"])
    lines.append("")
    for t in schema.tables:
        lines.append(f"### 4.{schema.tables.index(t) + 1} `{t.name}`")
        if t.comment:
            lines.append(f"> {t.comment}")
            lines.append("")
        lines.append(
            f"| {T['col_field']} | {T['col_type']} | {T['col_null']} | {T['col_pk']} | "
            f"{T['col_uniq']} | {T['col_default']} | {T['col_comment']} |"
        )
        lines.append("|------|------|------|------|------|--------|------|")
        for c in t.columns:
            lines.append(
                f"| `{c.name}` | {c.data_type} | {T['yes'] if c.nullable else T['no']} | "
                f"{'✓' if c.is_primary_key else ''} | {'✓' if c.is_unique else ''} | "
                f"{c.default_value or ''} | {c.comment or ''} |"
            )
        lines.append("")

    # 5 规范化分析
    lines.append(T["norm"])
    lines.append("")
    lines.append(T["norm_note"])
    lines.append("")
    for t in schema.tables:
        pk = T["join"].join(c.name for c in t.columns if c.is_primary_key)
        lines.append(T["norm_line"].format(table=t.name, pk=pk or T["no_pk"]))
    lines.append("")

    # 6 关系说明
    lines.append(T["rels"])
    lines.append("")
    lines.append(f"| {T['rel_cols'][0]} | {T['rel_cols'][1]} | {T['rel_cols'][2]} | {T['rel_cols'][3]} | {T['rel_cols'][4]} |")
    lines.append("|-----------|-------------|------|------|------|")
    for r in relationships:
        lines.append(
            f"| `{r.get('source_table')}.{r.get('source_column')}` | "
            f"`{r.get('target_table')}.{r.get('target_column')}` | "
            f"{r.get('cardinality') or ''} | {r.get('source_type') or ''} | {r.get('status') or ''} |"
        )
    lines.append("")

    # 7 数据字典
    lines.append(T["dict"])
    lines.append("")
    lines.append(T["dict_note"])
    lines.append("")

    # 8 总结
    lines.append(T["summary"])
    lines.append("")
    lines.append(
        T["summary_text"].format(
            tables=len(schema.tables),
            fields=sum(len(t.columns) for t in schema.tables),
            rels=len(relationships),
        )
    )
    lines.append("")
    return "\n".join(lines)


def build_word(
    project_name: str,
    project_description: str | None,
    schema: ParsedSchema,
    relationships: list[dict],
    concept_model: dict | None = None,
    author: str = "",
    student_id: str = "",
    lang: str = "zh",
) -> bytes:
    from docx import Document

    T = _texts(lang)
    doc = Document()
    doc.add_heading(T["doc_title"].format(project=project_name), 0)
    if author or student_id:
        doc.add_paragraph(T["author_line"].format(author=author or "-", student=student_id or "-"))
    doc.add_paragraph(T["generated"].format(now=datetime.now().strftime("%Y-%m-%d %H:%M:%S")))

    entities = concept_model.get("entities", []) if concept_model else []
    relations = concept_model.get("relations", []) if concept_model else []

    doc.add_heading(T["overview"].lstrip("# "), level=1)
    if project_description:
        doc.add_paragraph(project_description)
    doc.add_paragraph(
        f"{T['db_name'].format(name=schema.database_name)}; {T['table_count'].format(n=len(schema.tables))}; "
        f"{T['rel_count'].format(n=len(relationships))}"
    )

    doc.add_heading(T["req"].lstrip("# "), level=1)
    doc.add_paragraph(T["req_note"])
    for e in entities:
        attrs = T["join"].join(a["name"] for a in (e.get("attributes") or [])[:8])
        doc.add_paragraph(f"Entity {e['name']}{T['colon']}{attrs}", style="List Bullet")
    for r in relations:
        src = _entity_name(entities, r.get("source_entity") or "")
        tgt = _entity_name(entities, r.get("target_entity") or "")
        doc.add_paragraph(
            f"Relationship {r['name']}{T['colon']}{src}{T['paren'].format(v=r.get('source_card'))} "
            f"— {tgt}{T['paren'].format(v=r.get('target_card'))}",
            style="List Bullet",
        )

    doc.add_heading(T["concept"].lstrip("# "), level=1)
    doc.add_paragraph(T["concept_note"])
    for e in entities:
        doc.add_heading(e["name"], level=2)
        for a in e.get("attributes") or []:
            tag = T["pk_tag"] if a.get("is_pk") else ""
            doc.add_paragraph(f"{a['name']}{tag}", style="List Bullet")

    doc.add_heading(T["logical"].lstrip("# "), level=1)
    for t in schema.tables:
        doc.add_heading(t.name, level=2)
        if t.comment:
            doc.add_paragraph(t.comment)
        table = doc.add_table(rows=1, cols=7)
        table.style = "Light Grid Accent 1"
        for i, text in enumerate(
            [T["col_field"], T["col_type"], T["col_null"], T["col_pk"], T["col_uniq"], T["col_default"], T["col_comment"]]
        ):
            table.rows[0].cells[i].text = text
        for c in t.columns:
            cells = table.add_row().cells
            cells[0].text = c.name
            cells[1].text = c.data_type
            cells[2].text = T["yes"] if c.nullable else T["no"]
            cells[3].text = "✓" if c.is_primary_key else ""
            cells[4].text = "✓" if c.is_unique else ""
            cells[5].text = c.default_value or ""
            cells[6].text = c.comment or ""
        doc.add_paragraph()

    doc.add_heading(T["norm"].lstrip("# "), level=1)
    doc.add_paragraph(T["norm_note"])
    for t in schema.tables:
        pk = T["join"].join(c.name for c in t.columns if c.is_primary_key)
        doc.add_paragraph(T["norm_line"].format(table=t.name, pk=pk or T["no_pk"]), style="List Bullet")

    doc.add_heading(T["rels"].lstrip("# "), level=1)
    table = doc.add_table(rows=1, cols=4)
    table.style = "Light Grid Accent 1"
    for i, text in enumerate(T["rel_cols"][:4]):
        table.rows[0].cells[i].text = text
    for r in relationships:
        cells = table.add_row().cells
        cells[0].text = f"{r.get('source_table')}.{r.get('source_column')}"
        cells[1].text = f"{r.get('target_table')}.{r.get('target_column')}"
        cells[2].text = r.get("cardinality") or ""
        cells[3].text = r.get("source_type") or ""

    doc.add_heading(T["summary"].lstrip("# "), level=1)
    doc.add_paragraph(
        T["summary_text"].format(
            tables=len(schema.tables),
            fields=sum(len(t.columns) for t in schema.tables),
            rels=len(relationships),
        )
    )

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def build_pdf(
    project_name: str,
    project_description: str | None,
    schema: ParsedSchema,
    relationships: list[dict],
    concept_model: dict | None = None,
    author: str = "",
    student_id: str = "",
) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=14 * mm, rightMargin=14 * mm)
    styles = getSampleStyleSheet()
    body = ParagraphStyle("BodyCN", parent=styles["Normal"], fontSize=9, leading=13)
    h1 = ParagraphStyle("H1CN", parent=styles["Heading1"], fontSize=15, spaceBefore=8, spaceAfter=6)
    h2 = ParagraphStyle("H2CN", parent=styles["Heading2"], fontSize=12, spaceBefore=6, spaceAfter=4)
    story = [Paragraph(f"{project_name} 数据库设计说明书", styles["Title"])]
    story.append(Paragraph(f"作者：{author or '-'}　学号：{student_id or '-'}", body))
    story.append(Paragraph(f"自动生成于 {datetime.now().strftime('%Y-%m-%d %H:%M:%S')} · AI Database Architect", body))

    story.append(Paragraph("1. 项目概述", h1))
    if project_description:
        story.append(Paragraph(project_description, body))
    story.append(Paragraph(f"数据库名称：{schema.database_name}；表数量：{len(schema.tables)}；关系数量：{len(relationships)}", body))

    story.append(Paragraph("2. 需求分析", h1))
    story.append(Paragraph("说明：请根据实际业务需求补充需求描述。以下为从数据库结构中识别出的实体与联系概要。", body))
    entities = concept_model.get("entities", []) if concept_model else []
    for e in entities:
        attrs = "、".join(a["name"] for a in (e.get("attributes") or [])[:8])
        story.append(Paragraph(f"实体 {e['name']}：{attrs}", body))

    story.append(Paragraph("3. 概念结构设计", h1))
    story.append(Paragraph("采用 Chen 记法 ER 概念模型描述业务实体、属性与联系。", body))
    for e in entities:
        story.append(Paragraph(e["name"], h2))
        for a in e.get("attributes") or []:
            tag = "（主键）" if a.get("is_pk") else ""
            story.append(Paragraph(f"· {a['name']}{tag}", body))

    story.append(Paragraph("4. 逻辑结构设计", h1))
    for t in schema.tables:
        story.append(Paragraph(t.name, h2))
        rows = [["字段", "类型", "可空", "主键", "唯一", "默认值", "注释"]]
        for c in t.columns:
            rows.append(
                [c.name, c.data_type, "是" if c.nullable else "否",
                 "✓" if c.is_primary_key else "", "✓" if c.is_unique else "",
                 c.default_value or "", c.comment or ""]
            )
        tbl = Table(rows, repeatRows=1, colWidths=[30 * mm, 24 * mm, 10 * mm, 10 * mm, 10 * mm, 22 * mm, 44 * mm])
        tbl.setStyle(
            TableStyle(
                [
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DBEAFE")),
                    ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#CBD5E1")),
                    ("FONTSIZE", (0, 0), (-1, -1), 7),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ]
            )
        )
        story.append(tbl)
        story.append(Spacer(1, 6))

    story.append(Paragraph("5. 规范化分析", h1))
    story.append(Paragraph("提示：可结合『架构评审』中的 AI 意见补充范式分析章节。", body))
    for t in schema.tables:
        pk = "、".join(c.name for c in t.columns if c.is_primary_key)
        story.append(Paragraph(f"· {t.name}：主键 {pk or '无'}；满足 1NF。", body))

    story.append(Paragraph("6. 关系说明", h1))
    rows = [["源表.字段", "目标表.字段", "基数", "类型"]]
    for r in relationships:
        rows.append(
            [
                f"{r.get('source_table')}.{r.get('source_column')}",
                f"{r.get('target_table')}.{r.get('target_column')}",
                r.get("cardinality") or "",
                r.get("source_type") or "",
            ]
        )
    tbl = Table(rows, repeatRows=1, colWidths=[52 * mm, 52 * mm, 30 * mm, 30 * mm])
    tbl.setStyle(
        TableStyle(
            [
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DBEAFE")),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#CBD5E1")),
                ("FONTSIZE", (0, 0), (-1, -1), 7.5),
            ]
        )
    )
    story.append(tbl)

    story.append(Paragraph("7. 总结", h1))
    story.append(
        Paragraph(
            f"本系统数据库共设计 {len(schema.tables)} 张表、{sum(len(t.columns) for t in schema.tables)} 个字段、"
            f"{len(relationships)} 条关系，覆盖了业务的核心数据需求，并遵循数据库规范化原则。",
            body,
        )
    )
    doc.build(story)
    return buf.getvalue()
