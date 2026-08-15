"""Data dictionary export - Markdown / Excel / Word / PDF (feature 3)."""
from __future__ import annotations

import io
from datetime import datetime

from app.services.schema_parser import ParsedSchema


def _merged_comments(
    schema: ParsedSchema,
    accepted: list[dict],
) -> tuple[dict[str, str], dict[tuple[str, str], str]]:
    """Merge existing schema comments with accepted AI suggestions."""
    table_comments = {t.name: (t.comment or "") for t in schema.tables}
    column_comments = {
        (t.name, c.name): (c.comment or "") for t in schema.tables for c in t.columns
    }
    for row in accepted or []:
        if row.get("target_type") == "table":
            table_comments[str(row["table_name"])] = str(row["suggested_comment"])
        else:
            column_comments[(str(row["table_name"]), str(row["column_name"] or ""))] = str(
                row["suggested_comment"]
            )
    return table_comments, column_comments


def _statistics(schema: ParsedSchema, tc, cc) -> dict:
    tables_missing = sum(1 for t in schema.tables if not tc.get(t.name, "").strip())
    cols_missing = sum(
        1 for t in schema.tables for c in t.columns if not cc.get((t.name, c.name), "").strip()
    )
    return {
        "tables": len(schema.tables),
        "columns": sum(len(t.columns) for t in schema.tables),
        "tables_missing": tables_missing,
        "columns_missing": cols_missing,
    }


def export_markdown(schema: ParsedSchema, accepted: list[dict]) -> str:
    tc, cc = _merged_comments(schema, accepted)
    stats = _statistics(schema, tc, cc)
    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    lines = [
        "# 数据字典",
        "",
        f"> 自动生成于 {now} · AI Database Architect",
        "",
        f"- 表数量：{stats['tables']} · 字段数量：{stats['columns']}",
        f"- 缺表注释：{stats['tables_missing']} · 缺字段注释：{stats['columns_missing']}",
        "",
    ]
    for t in schema.tables:
        lines.append(f"## {t.name}")
        lines.append("")
        if tc.get(t.name):
            lines.append(f"> {tc[t.name]}")
            lines.append("")
        lines.append("| 字段 | 类型 | 可空 | 主键 | 唯一 | 默认值 | 注释 |")
        lines.append("|------|------|------|------|------|--------|------|")
        for c in t.columns:
            lines.append(
                f"| `{c.name}` | {c.data_type} | {'是' if c.nullable else '否'} | "
                f"{'✓' if c.is_primary_key else ''} | {'✓' if c.is_unique else ''} | "
                f"{c.default_value or ''} | {cc.get((t.name, c.name), '') or ''} |"
            )
        lines.append("")
    return "\n".join(lines)


def export_excel(schema: ParsedSchema, accepted: list[dict]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    tc, cc = _merged_comments(schema, accepted)
    wb = Workbook()

    # Overview sheet
    ws = wb.active
    ws.title = "概览"
    ws.append(["表名", "表注释", "字段数", "主键字段"])
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DBEAFE")
    for t in schema.tables:
        ws.append(
            [
                t.name,
                tc.get(t.name, ""),
                len(t.columns),
                "、".join(c.name for c in t.columns if c.is_primary_key),
            ]
        )
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 40
    ws.column_dimensions["C"].width = 10
    ws.column_dimensions["D"].width = 30

    # One sheet per table
    for t in schema.tables:
        ws = wb.create_sheet(t.name[:31] or "sheet")
        ws.append(["字段名", "类型", "可空", "主键", "唯一", "默认值", "注释"])
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="DBEAFE")
        for c in t.columns:
            ws.append(
                [
                    c.name,
                    c.data_type + (f"({c.length})" if c.length else ""),
                    "是" if c.nullable else "否",
                    "✓" if c.is_primary_key else "",
                    "✓" if c.is_unique else "",
                    c.default_value or "",
                    cc.get((t.name, c.name), "") or "",
                ]
            )
        for row in ws.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="center")
        widths = [22, 18, 8, 8, 8, 16, 46]
        for idx, w in enumerate(widths, start=1):
            ws.column_dimensions[chr(64 + idx)].width = w

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_word(schema: ParsedSchema, accepted: list[dict]) -> bytes:
    from docx import Document
    from docx.shared import Pt

    tc, cc = _merged_comments(schema, accepted)
    doc = Document()
    doc.add_heading("数据字典", 0)
    doc.add_paragraph(f"自动生成于 {datetime.now().strftime('%Y-%m-%d %H:%M:%S')} · AI Database Architect")

    for t in schema.tables:
        doc.add_heading(t.name, level=1)
        if tc.get(t.name):
            doc.add_paragraph(tc[t.name])
        table = doc.add_table(rows=1, cols=7)
        table.style = "Light Grid Accent 1"
        header = table.rows[0].cells
        for i, text in enumerate(["字段名", "类型", "可空", "主键", "唯一", "默认值", "注释"]):
            header[i].text = text
        for c in t.columns:
            cells = table.add_row().cells
            cells[0].text = c.name
            cells[1].text = c.data_type + (f"({c.length})" if c.length else "")
            cells[2].text = "是" if c.nullable else "否"
            cells[3].text = "✓" if c.is_primary_key else ""
            cells[4].text = "✓" if c.is_unique else ""
            cells[5].text = c.default_value or ""
            cells[6].text = cc.get((t.name, c.name), "") or ""
        doc.add_paragraph()

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def export_pdf(schema: ParsedSchema, accepted: list[dict]) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import (
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    tc, cc = _merged_comments(schema, accepted)
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=14 * mm, rightMargin=14 * mm)
    styles = getSampleStyleSheet()
    body = ParagraphStyle("BodyCN", parent=styles["Normal"], fontSize=9, leading=13)
    h1 = ParagraphStyle("H1CN", parent=styles["Heading1"], fontSize=16, spaceAfter=8)
    h2 = ParagraphStyle("H2CN", parent=styles["Heading2"], fontSize=13, spaceBefore=10, spaceAfter=4)

    story = [Paragraph("数据字典", h1)]
    story.append(
        Paragraph(
            f"自动生成于 {datetime.now().strftime('%Y-%m-%d %H:%M:%S')} · AI Database Architect",
            body,
        )
    )
    story.append(Spacer(1, 6))

    for t in schema.tables:
        story.append(Paragraph(t.name, h2))
        if tc.get(t.name):
            story.append(Paragraph(tc[t.name], body))
        rows = [["字段名", "类型", "可空", "主键", "唯一", "默认值", "注释"]]
        for c in t.columns:
            rows.append(
                [
                    c.name,
                    c.data_type + (f"({c.length})" if c.length else ""),
                    "是" if c.nullable else "否",
                    "✓" if c.is_primary_key else "",
                    "✓" if c.is_unique else "",
                    c.default_value or "",
                    cc.get((t.name, c.name), "") or "",
                ]
            )
        tbl = Table(rows, repeatRows=1, colWidths=[32 * mm, 26 * mm, 12 * mm, 12 * mm, 12 * mm, 24 * mm, 44 * mm])
        tbl.setStyle(
            TableStyle(
                [
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DBEAFE")),
                    ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#CBD5E1")),
                    ("FONTSIZE", (0, 0), (-1, -1), 7.5),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
                ]
            )
        )
        story.append(tbl)
        story.append(Spacer(1, 8))

    doc.build(story)
    return buf.getvalue()
